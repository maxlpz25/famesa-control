from flask import Flask, render_template, request, redirect, url_for, session, jsonify, flash, send_file
from werkzeug.security import generate_password_hash, check_password_hash
import sqlite3
import os
import io
import csv
from datetime import datetime

app = Flask(__name__)
app.secret_key = "famesa_explosivos_planta_2026_@hyperpro"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(BASE_DIR, "famesa_multisede.db")

SEDES_VALIDAS = ["Puente Piedra", "Chancay", "La Joya"]

def get_db():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db()
    c = conn.cursor()

    # 1. Tabla Usuarios (con nuevos permisos de auditoría y edición)
    c.execute('''
    CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        nombre_completo TEXT NOT NULL,
        rol TEXT NOT NULL,
        sede_asignada TEXT NOT NULL DEFAULT 'TODAS',
        puede_registrar INTEGER DEFAULT 1,
        puede_editar INTEGER DEFAULT 0,
        puede_borrar INTEGER DEFAULT 0,
        puede_gestionar_categorias INTEGER DEFAULT 0,
        puede_exportar_excel INTEGER DEFAULT 1,
        puede_crear_usuarios INTEGER DEFAULT 0,
        creado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')

    # 2. Tabla Categorías Dinámicas
    c.execute('''
    CREATE TABLE IF NOT EXISTS categorias (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT UNIQUE NOT NULL,
        aplica_comision INTEGER DEFAULT 0,
        color_tag TEXT DEFAULT '#E30613',
        creado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')

    # 3. Tabla Movimientos Operativos
    c.execute('''
    CREATE TABLE IF NOT EXISTS movimientos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        sede TEXT NOT NULL,
        fecha TEXT NOT NULL,
        hora TEXT NOT NULL,
        categoria TEXT NOT NULL,
        tipo_movimiento TEXT NOT NULL,
        cantidad INTEGER NOT NULL DEFAULT 1,
        empresa TEXT DEFAULT '-',
        placa TEXT DEFAULT '-',
        supervisor TEXT DEFAULT '-',
        motivo TEXT DEFAULT '-',
        registrado_por TEXT,
        creado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')

    # 4. Tabla de Auditoría / Rastro de Cambios (Solo visible para Administrador)
    c.execute('''
    CREATE TABLE IF NOT EXISTS auditoria (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        usuario TEXT NOT NULL,
        accion TEXT NOT NULL,
        detalles TEXT NOT NULL,
        fecha_hora TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )
    ''')

    # Categorías iniciales
    cats_default = [
        ("Personal Propio Famesa", 1, "#E30613"),
        ("Contratas / Servicios", 0, "#059669"),
        ("Transportistas / Carga", 0, "#2563EB"),
        ("Custodios / Seguridad", 0, "#7C3AED"),
        ("Policías / DINOES", 0, "#D97706"),
        ("Visitas Técnicas / Terceros", 0, "#0891B2")
    ]
    for nom, com, col in cats_default:
        c.execute("INSERT OR IGNORE INTO categorias (nombre, aplica_comision, color_tag) VALUES (?, ?, ?)", (nom, com, col))

    # Administrador por defecto
    c.execute("SELECT * FROM usuarios WHERE username = 'admin'")
    if not c.fetchone():
        admin_hash = generate_password_hash("admin123")
        c.execute('''
            INSERT INTO usuarios (username, password_hash, nombre_completo, rol, sede_asignada,
                                  puede_registrar, puede_editar, puede_borrar, puede_gestionar_categorias,
                                  puede_exportar_excel, puede_crear_usuarios)
            VALUES (?, ?, ?, 'superadmin', 'TODAS', 1, 1, 1, 1, 1, 1)
        ''', ('admin', admin_hash, 'Administrador General FAMESA'))

    conn.commit()
    conn.close()

@app.route("/")
def index():
    return render_template("dashboard.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = request.form.get("username", "").strip()
        p = request.form.get("password", "").strip()
        conn = get_db()
        user = conn.execute("SELECT * FROM usuarios WHERE username = ?", (u,)).fetchone()
        conn.close()

        if user and check_password_hash(user["password_hash"], p):
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["nombre"] = user["nombre_completo"]
            session["rol"] = user["rol"]
            session["sede_asignada"] = user["sede_asignada"]
            session["puede_registrar"] = bool(user["puede_registrar"])
            session["puede_editar"] = bool(user["puede_editar"])
            session["puede_borrar"] = bool(user["puede_borrar"])
            session["puede_gestionar_categorias"] = bool(user["puede_gestionar_categorias"])
            session["puede_exportar_excel"] = bool(user["puede_exportar_excel"])
            session["puede_crear_usuarios"] = bool(user["puede_crear_usuarios"])
            return redirect(url_for("panel_operativo"))
        else:
            flash("Credenciales incorrectas", "error")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("index"))

@app.route("/panel")
def panel_operativo():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("panel_operativo.html")

@app.route("/usuarios")
def gestion_usuarios():
    if "user_id" not in session or not session.get("puede_crear_usuarios"):
        flash("Acceso denegado.", "error")
        return redirect(url_for("panel_operativo"))
    conn = get_db()
    users = conn.execute("SELECT * FROM usuarios ORDER BY id ASC").fetchall()
    
    # Obtener el registro de auditoría exclusivo para el Administrador
    audits = conn.execute("SELECT * FROM auditoria ORDER BY id DESC LIMIT 100").fetchall()
    conn.close()
    return render_template("usuarios.html", usuarios=users, sedes=SEDES_VALIDAS, auditoria=audits)

@app.route("/api/mi_perfil", methods=["PUT"])
def api_mi_perfil():
    if "user_id" not in session:
        return jsonify({"error": "No autenticado"}), 401
    d = request.json
    nuevo_username = d.get("username", "").strip()
    nueva_password = d.get("password", "").strip()
    user_id = session["user_id"]
    if not nuevo_username:
        return jsonify({"error": "Usuario requerido"}), 400
    conn = get_db()
    c = conn.cursor()
    try:
        if nueva_password:
            nueva_hash = generate_password_hash(nueva_password)
            c.execute("UPDATE usuarios SET username = ?, password_hash = ? WHERE id = ?", (nuevo_username, nueva_hash, user_id))
        else:
            c.execute("UPDATE usuarios SET username = ? WHERE id = ?", (nuevo_username, user_id))
        conn.commit()
        session["username"] = nuevo_username
        conn.close()
        return jsonify({"success": True})
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "El usuario ya existe"}), 400

# GESTIÓN Y EDICIÓN DE CATEGORÍAS
@app.route("/api/categorias", methods=["GET", "POST", "PUT", "DELETE"])
def api_categorias():
    conn = get_db()
    c = conn.cursor()

    if request.method == "POST":
        if "user_id" not in session or not session.get("puede_gestionar_categorias"):
            conn.close()
            return jsonify({"error": "No autorizado"}), 403
        data = request.json
        nombre = data.get("nombre", "").strip()
        aplica_com = 1 if data.get("aplica_comision") else 0
        color = data.get("color_tag", "#E30613")
        if not nombre:
            conn.close()
            return jsonify({"error": "Nombre requerido"}), 400
        try:
            c.execute("INSERT INTO categorias (nombre, aplica_comision, color_tag) VALUES (?, ?, ?)", (nombre, aplica_com, color))
            c.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
                      (session.get("username"), "CREAR CATEGORÍA", f"Se creó la categoría: {nombre}"))
            conn.commit()
            conn.close()
            return jsonify({"success": True})
        except sqlite3.IntegrityError:
            conn.close()
            return jsonify({"error": "Ya existe"}), 400

    elif request.method == "PUT":
        if "user_id" not in session or not session.get("puede_gestionar_categorias"):
            conn.close()
            return jsonify({"error": "No autorizado"}), 403
        data = request.json
        cat_id = data.get("id")
        nuevo_nombre = data.get("nombre", "").strip()
        aplica_com = 1 if data.get("aplica_comision") else 0
        color = data.get("color_tag", "#E30613")

        if not cat_id or not nuevo_nombre:
            conn.close()
            return jsonify({"error": "Datos incompletos"}), 400

        cat_vieja = c.execute("SELECT nombre FROM categorias WHERE id = ?", (cat_id,)).fetchone()
        nombre_viejo = cat_vieja["nombre"] if cat_vieja else ""

        try:
            c.execute("UPDATE categorias SET nombre = ?, aplica_comision = ?, color_tag = ? WHERE id = ?", 
                      (nuevo_nombre, aplica_com, color, cat_id))
            # Actualizar también en movimientos para mantener consistencia
            c.execute("UPDATE movimientos SET categoria = ? WHERE categoria = ?", (nuevo_nombre, nombre_viejo))
            c.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
                      (session.get("username"), "EDITAR CATEGORÍA", f"Se editó la categoría '{nombre_viejo}' a '{nuevo_nombre}'"))
            conn.commit()
            conn.close()
            return jsonify({"success": True})
        except sqlite3.IntegrityError:
            conn.close()
            return jsonify({"error": "Ese nombre de categoría ya está en uso"}), 400

    elif request.method == "DELETE":
        if "user_id" not in session or not session.get("puede_gestionar_categorias"):
            conn.close()
            return jsonify({"error": "No autorizado"}), 403
        cat_id = request.args.get("id")
        cat_row = c.execute("SELECT nombre FROM categorias WHERE id = ?", (cat_id,)).fetchone()
        if cat_row:
            cat_nom = cat_row["nombre"]
            c.execute("DELETE FROM movimientos WHERE categoria = ?", (cat_nom,))
            c.execute("DELETE FROM categorias WHERE id = ?", (cat_id,))
            c.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
                      (session.get("username"), "ELIMINAR CATEGORÍA", f"Se eliminó la categoría: {cat_nom}"))
            conn.commit()
        conn.close()
        return jsonify({"success": True})

    cats = c.execute("SELECT * FROM categorias ORDER BY id ASC").fetchall()
    conn.close()
    return jsonify([dict(row) for row in cats])

@app.route("/api/resumen")
def api_resumen():
    sede = request.args.get("sede", "Puente Piedra")
    fecha = request.args.get("fecha", datetime.now().strftime("%Y-%m-%d"))
    conn = get_db()
    c = conn.cursor()
    cats_db = c.execute("SELECT * FROM categorias ORDER BY id ASC").fetchall()
    cats_map = {row["nombre"]: dict(row) for row in cats_db}

    if sede == "TODAS":
        movs = c.execute("SELECT * FROM movimientos WHERE fecha = ? ORDER BY id DESC", (fecha,)).fetchall()
    else:
        movs = c.execute("SELECT * FROM movimientos WHERE sede = ? AND fecha = ? ORDER BY id DESC", (sede, fecha)).fetchall()
    conn.close()

    metricas = {}
    for nom in cats_map:
        metricas[nom] = {
            "ingresos": 0, "salidas": 0, "com_sal": 0, "com_ret": 0,
            "com_activa": 0, "presentes": 0,
            "aplica_comision": cats_map[nom]["aplica_comision"],
            "color_tag": cats_map[nom]["color_tag"]
        }

    for m in movs:
        cat = m["categoria"]
        if cat in metricas:
            t = m["tipo_movimiento"]
            cant = int(m["cantidad"] or 0)
            if t == "INGRESO TURNO":
                metricas[cat]["ingresos"] += cant
            elif t == "SALIDA TURNO":
                metricas[cat]["salidas"] += cant
            elif t == "SALIDA COMISIÓN":
                metricas[cat]["com_sal"] += cant
            elif t == "RETORNO COMISIÓN":
                metricas[cat]["com_ret"] += cant

    tot_ing = tot_sal = tot_com_s = tot_com_r = tot_com_act = tot_presentes = 0
    tabla_resumen = []
    tarjetas_dinamicas = []

    for nom, item in metricas.items():
        if item["aplica_comision"] == 1:
            item["com_activa"] = max(0, item["com_sal"] - item["com_ret"])
            item["presentes"] = max(0, (item["ingresos"] - item["salidas"]) - item["com_activa"])
        else:
            item["com_activa"] = 0
            item["presentes"] = max(0, item["ingresos"] - item["salidas"])

        tot_ing += item["ingresos"]
        tot_sal += item["salidas"]
        tot_com_s += item["com_sal"]
        tot_com_r += item["com_ret"]
        tot_com_act += item["com_activa"]
        tot_presentes += item["presentes"]

        tarjetas_dinamicas.append({
            "categoria": nom,
            "presentes": item["presentes"],
            "color_tag": item["color_tag"],
            "ingresos": item["ingresos"],
            "salidas": item["salidas"],
            "com_activa": item["com_activa"] if item["aplica_comision"] else 0
        })

        tabla_resumen.append({
            "categoria": nom,
            "color_tag": item["color_tag"],
            "ingresos": item["ingresos"],
            "salidas": item["salidas"],
            "com_sal": item["com_sal"] if item["aplica_comision"] else "-",
            "com_ret": item["com_ret"] if item["aplica_comision"] else "-",
            "com_activa": item["com_activa"] if item["aplica_comision"] else "-",
            "presentes": item["presentes"]
        })

    return jsonify({
        "sede": sede,
        "fecha": fecha,
        "kpis": {
            "total_planta": tot_presentes,
            "total_ingresos": tot_ing,
            "total_salidas": tot_sal
        },
        "tarjetas": tarjetas_dinamicas,
        "tabla": tabla_resumen,
        "totales_pie": {
            "ingresos": tot_ing,
            "salidas": tot_sal,
            "com_sal": tot_com_s,
            "com_ret": tot_com_r,
            "com_activa": tot_com_act,
            "presentes": tot_presentes
        }
    })

# REGISTRO Y EDICIÓN DE MOVIMIENTOS (CON RASTRO DE AUDITORÍA)
@app.route("/api/movimientos", methods=["GET", "POST", "PUT"])
def api_movimientos():
    conn = get_db()
    c = conn.cursor()

    if request.method == "POST":
        if "user_id" not in session or not session.get("puede_registrar"):
            conn.close()
            return jsonify({"error": "Sin permisos"}), 403
        data = request.json
        sede = data.get("sede")
        sede_user = session.get("sede_asignada")
        if sede_user != "TODAS" and sede_user != sede:
            conn.close()
            return jsonify({"error": "Acceso denegado"}), 403

        c.execute('''
            INSERT INTO movimientos (sede, fecha, hora, categoria, tipo_movimiento, cantidad, empresa, placa, supervisor, motivo, registrado_por)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            sede,
            data.get("fecha") or datetime.now().strftime("%Y-%m-%d"),
            data.get("hora") or datetime.now().strftime("%H:%M:%S"),
            data.get("categoria"),
            data.get("tipo_movimiento"),
            int(data.get("cantidad") or 1),
            (data.get("empresa") or "").strip() or "-",
            (data.get("placa") or "").strip() or "-",
            (data.get("supervisor") or "").strip() or "-",
            (data.get("motivo") or "").strip() or "-",
            session.get("username")
        ))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    elif request.method == "PUT":
        if "user_id" not in session or not session.get("puede_editar"):
            conn.close()
            return jsonify({"error": "No autorizado para editar registros"}), 403
        
        data = request.json
        mov_id = data.get("id")
        if not mov_id:
            conn.close()
            return jsonify({"error": "ID requerido"}), 400

        c.execute('''
            UPDATE movimientos 
            SET categoria = ?, tipo_movimiento = ?, cantidad = ?, empresa = ?, placa = ?, supervisor = ?, motivo = ?
            WHERE id = ?
        ''', (
            data.get("categoria"),
            data.get("tipo_movimiento"),
            int(data.get("cantidad") or 1),
            (data.get("empresa") or "").strip() or "-",
            (data.get("placa") or "").strip() or "-",
            (data.get("supervisor") or "").strip() or "-",
            (data.get("motivo") or "").strip() or "-",
            mov_id
        ))
        
        # Registrar auditoría de edición de movimiento
        c.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
                  (session.get("username"), "EDITAR MOVIMIENTO", f"Se editó el registro ID #{mov_id}"))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    sede = request.args.get("sede", "Puente Piedra")
    fecha = request.args.get("fecha")
    mes = request.args.get("mes")
    sql = "SELECT * FROM movimientos WHERE 1=1"
    params = []
    if sede != "TODAS":
        sql += " AND sede = ?"
        params.append(sede)
    if mes:
        sql += " AND fecha LIKE ?"
        params.append(f"{mes}%")
    elif fecha:
        sql += " AND fecha = ?"
        params.append(fecha)
    sql += " ORDER BY id DESC"
    rows = c.execute(sql, params).fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows])

@app.route("/api/movimientos/<int:id>", methods=["DELETE"])
def api_movimiento_del(id):
    if "user_id" not in session or not session.get("puede_borrar"):
        return jsonify({"error": "Sin permisos"}), 403
    conn = get_db()
    c = conn.cursor()
    c.execute("DELETE FROM movimientos WHERE id = ?", (id,))
    c.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
              (session.get("username"), "ELIMINAR MOVIMIENTO", f"Se eliminó el registro ID #{id}"))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

# EXPORTACIÓN EXCEL EN MEMORIA RAM
@app.route("/exportar_excel")
def exportar_excel():
    if "user_id" not in session or not session.get("puede_exportar_excel"):
        return "Acceso denegado", 403

    sede = request.args.get("sede", "TODAS").strip()
    mes = request.args.get("mes", "").strip()
    fecha = request.args.get("fecha", "").strip()

    if not mes and not fecha:
        mes = datetime.now().strftime("%Y-%m")

    conn = get_db()
    c = conn.cursor()
    sql = "SELECT * FROM movimientos WHERE 1=1"
    params = []
    if sede and sede != "TODAS":
        sql += " AND sede = ?"
        params.append(sede)
    if fecha:
        sql += " AND fecha = ?"
        params.append(fecha)
    elif mes:
        sql += " AND fecha LIKE ?"
        params.append(f"{mes}%")
    sql += " ORDER BY fecha DESC, hora DESC"
    rows = c.execute(sql, params).fetchall()
    conn.close()

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control_Personal"
    ws.views.sheetView[0].showGridLines = True

    famesa_red = PatternFill("solid", fgColor="E30613")
    navy_header = PatternFill("solid", fgColor="0F172A")
    sub_fill = PatternFill("solid", fgColor="1E293B")
    zebra_light = PatternFill("solid", fgColor="F8FAFC")
    zebra_white = PatternFill("solid", fgColor="FFFFFF")

    font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_meta_lbl = Font(name="Segoe UI", size=9, bold=True, color="475569")
    font_meta_val = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_data = Font(name="Segoe UI", size=9, color="1E293B")
    font_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")

    thin_gray = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    ws.merge_cells("A1:L2")
    ws["A1"] = "FAMESA EXPLOSIVOS - REPORTE GERENCIAL DE CONTROL DE PERSONAL"
    ws["A1"].font = font_title
    ws["A1"].fill = famesa_red
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[3].height = 20
    meta_info = [
        ("A3", "B3", "SEDE / PLANTA:", sede if sede else "TODAS LAS SEDES"),
        ("D3", "E3", "PERIODO:", fecha if fecha else (f"Mes de {mes}")),
        ("G3", "H3", "EMISIÓN:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("J3", "K3", "REGISTROS:", len(rows))
    ]
    for c_lbl, c_val, lbl, val in meta_info:
        ws[c_lbl] = lbl
        ws[c_lbl].font = font_meta_lbl
        ws[c_lbl].alignment = Alignment(horizontal="right", vertical="center")
        ws[c_val] = val
        ws[c_val].font = font_meta_val
        ws[c_val].alignment = Alignment(horizontal="left", vertical="center")

    headers = [
        "ID", "SEDE", "FECHA", "HORA (24H)", "CATEGORÍA / ESTRUCTURA", 
        "TIPO DE MOVIMIENTO", "CANTIDAD", "EMPRESA / DETALLE", "PLACA / VEHÍCULO", 
        "SUPERVISOR DE GUARDIA", "OBSERVACIÓN", "REGISTRADO POR"
    ]
    ws.row_dimensions[5].height = 26
    for col_idx, h_text in enumerate(headers, 1):
        c_head = ws.cell(row=5, column=col_idx, value=h_text)
        c_head.font = font_header
        c_head.fill = navy_header
        c_head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c_head.border = border_all

    total_personas = 0
    for r_idx, r in enumerate(rows, 6):
        ws.row_dimensions[r_idx].height = 21
        fill_row = zebra_light if r_idx % 2 == 0 else zebra_white
        cant = int(r["cantidad"] or 0)
        total_personas += cant

        mov_tipo = r["tipo_movimiento"]
        tipo_color = "059669" if "INGRESO" in mov_tipo or "RETORNO" in mov_tipo else "DC2626"
        font_tipo = Font(name="Segoe UI", size=9, bold=True, color=tipo_color)

        ws.cell(row=r_idx, column=1, value=r["id"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=2, value=r["sede"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=3, value=r["fecha"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=4, value=r["hora"]).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=5, value=r["categoria"]).alignment = Alignment(horizontal="left", vertical="center")
        
        c_t = ws.cell(row=r_idx, column=6, value=mov_tipo)
        c_t.alignment = Alignment(horizontal="center", vertical="center")
        c_t.font = font_tipo

        c_cant = ws.cell(row=r_idx, column=7, value=cant)
        c_cant.alignment = Alignment(horizontal="right", vertical="center")
        c_cant.number_format = "#,##0"
        c_cant.font = font_bold

        ws.cell(row=r_idx, column=8, value=r["empresa"] or "-").alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=9, value=r["placa"] or "-").alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r_idx, column=10, value=r["supervisor"] or "-").alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=11, value=r["motivo"] or "-").alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r_idx, column=12, value=r["registrado_por"] or "-").alignment = Alignment(horizontal="center", vertical="center")

        for col_idx in range(1, 13):
            cell_item = ws.cell(row=r_idx, column=col_idx)
            cell_item.fill = fill_row
            cell_item.border = border_all
            if col_idx not in [6, 7]:
                cell_item.font = font_data

    last_row = 6 + len(rows)
    ws.row_dimensions[last_row].height = 24
    ws.merge_cells(f"A{last_row}:F{last_row}")
    ws[f"A{last_row}"] = "TOTAL CONSOLIDADO:"
    ws[f"A{last_row}"].font = font_header
    ws[f"A{last_row}"].fill = sub_fill
    ws[f"A{last_row}"].alignment = Alignment(horizontal="right", vertical="center")

    c_tot = ws.cell(row=last_row, column=7, value=total_personas)
    c_tot.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    c_tot.fill = famesa_red
    c_tot.alignment = Alignment(horizontal="right", vertical="center")
    c_tot.number_format = "#,##0"

    for col_idx in range(8, 13):
        ws.cell(row=last_row, column=col_idx, value="").fill = sub_fill

    for col_idx in range(1, 13):
        ws.cell(row=last_row, column=col_idx).border = border_all

    column_widths = [8, 16, 13, 14, 28, 22, 12, 28, 15, 22, 30, 16]
    for i, w in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename_clean = f"FAMESA_Reporte_{sede}_{mes or fecha}.xlsx".replace(" ", "_")
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename_clean
    )

@app.route("/api/usuarios", methods=["POST"])
def api_usuario_crear():
    if "user_id" not in session or not session.get("puede_crear_usuarios"):
        return jsonify({"error": "No autorizado"}), 403
    d = request.json
    try:
        conn = get_db()
        c = conn.cursor()
        c.execute('''
            INSERT INTO usuarios (username, password_hash, nombre_completo, rol, sede_asignada,
                                  puede_registrar, puede_editar, puede_borrar, puede_gestionar_categorias,
                                  puede_exportar_excel, puede_crear_usuarios)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            d.get("username"), generate_password_hash(d.get("password")), d.get("nombre"),
            d.get("rol", "operador"), d.get("sede_asignada", "Puente Piedra"),
            1 if d.get("puede_registrar") else 0, 1 if d.get("puede_editar") else 0,
            1 if d.get("puede_borrar") else 0, 1 if d.get("puede_gestionar_categorias") else 0,
            1 if d.get("puede_exportar_excel") else 0, 1 if d.get("rol") == "superadmin" else 0
        ))
        c.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
                  (session.get("username"), "CREAR USUARIO", f"Se creó el usuario: {d.get('username')}"))
        conn.commit()
        conn.close()
        return jsonify({"success": True})
    except:
        return jsonify({"error": "Error al crear"}), 400

@app.route("/api/usuarios/<int:id>", methods=["DELETE"])
def api_usuario_borrar(id):
    if "user_id" not in session or not session.get("puede_crear_usuarios"):
        return jsonify({"error": "No autorizado"}), 403
    conn = get_db()
    u = conn.execute("SELECT username FROM usuarios WHERE id = ?", (id,)).fetchone()
    username_del = u["username"] if u else str(id)
    conn.execute("DELETE FROM usuarios WHERE id = ?", (id,))
    conn.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
              (session.get("username"), "ELIMINAR USUARIO", f"Se eliminó al usuario: {username_del}"))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/usuarios/<int:id>", methods=["GET"])
def api_usuario_get(id):
    conn = get_db()
    u = conn.execute("SELECT * FROM usuarios WHERE id = ?", (id,)).fetchone()
    conn.close()
    return jsonify(dict(u))

@app.route("/api/usuarios/<int:id>/editar", methods=["PUT"])
def api_usuario_editar_completo(id):
    if "user_id" not in session or not session.get("puede_crear_usuarios"):
        return jsonify({"error": "No autorizado"}), 403
    d = request.json
    conn = get_db()
    c = conn.cursor()
    if d.get("password"):
        c.execute("UPDATE usuarios SET nombre_completo=?, username=?, password_hash=?, sede_asignada=?, rol=?, puede_registrar=?, puede_editar=?, puede_borrar=?, puede_gestionar_categorias=?, puede_exportar_excel=? WHERE id=?",
                  (d.get("nombre"), d.get("username"), generate_password_hash(d.get("password")), d.get("sede_asignada"), d.get("rol"), 
                   1 if d.get("puede_registrar") else 0, 1 if d.get("puede_editar") else 0, 1 if d.get("puede_borrar") else 0, 1 if d.get("puede_gestionar_categorias") else 0, 1 if d.get("puede_exportar_excel") else 0, id))
    else:
        c.execute("UPDATE usuarios SET nombre_completo=?, username=?, sede_asignada=?, rol=?, puede_registrar=?, puede_editar=?, puede_borrar=?, puede_gestionar_categorias=?, puede_exportar_excel=? WHERE id=?",
                  (d.get("nombre"), d.get("username"), d.get("sede_asignada"), d.get("rol"), 
                   1 if d.get("puede_registrar") else 0, 1 if d.get("puede_editar") else 0, 1 if d.get("puede_borrar") else 0, 1 if d.get("puede_gestionar_categorias") else 0, 1 if d.get("puede_exportar_excel") else 0, id))
    
    c.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (?, ?, ?)", 
              (session.get("username"), "EDITAR USUARIO", f"Se editaron permisos/datos del usuario ID #{id}"))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)